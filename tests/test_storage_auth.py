import io
import sqlite3
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from api import auth
from platform_core import config
from platform_core.store import ExcelStore, SCHEMA
from services.migration import migrate_legacy
from helpers import PlatformCase


class StorageAuthTests(PlatformCase):
    def test_requires_login_logout_and_origin_protection(self):
        anonymous = TestClient(self.application)
        self.assertEqual(anonymous.get("/api/agents").status_code, 401)
        self.assertEqual(
            anonymous.post("/api/chat", json={"message": "hello"}).status_code, 401
        )
        self.assertEqual(anonymous.get("/api/auth/me").json()["user"], None)
        self.assertEqual(
            self.client.post(
                "/api/agents",
                json={"name": "blocked", "domain": "HR", "type": "Chat"},
                headers={"Origin": "https://attacker.example"},
            ).status_code,
            403,
        )
        self.client.post("/api/auth/logout")
        self.assertEqual(self.client.get("/api/runs").status_code, 401)
        anonymous.close()

    def login_google(self, client, sub, email):
        identity = {
            "sub": sub,
            "email": email,
            "email_verified": True,
            "name": email.split("@")[0],
        }
        google = SimpleNamespace(
            authorize_access_token=AsyncMock(return_value={"userinfo": identity})
        )
        with (
            patch.object(config, "GOOGLE_CONFIGURED", True),
            patch.object(auth, "oauth", SimpleNamespace(google=google)),
        ):
            response = client.get("/api/auth/google/callback")
            self.assertEqual(response.status_code, 200)

    def test_google_users_are_isolated_in_agents_documents_reports_and_exports(self):
        alice = TestClient(self.application)
        bob = TestClient(self.application)
        self.login_google(alice, "alice", "alice@gmail.com")
        self.login_google(bob, "bob", "bob@gmail.com")
        self.assertEqual(
            alice.get("/api/auth/me").json()["user"]["email"], "alice@gmail.com"
        )
        agent = alice.post(
            "/api/agents/upload",
            data={"name": "Private policy", "domain": "HR", "type": "Chat"},
            files={"knowledge_files": ("private.txt", b"Vacation leave is 30 days.")},
        ).json()
        report = alice.post(
            "/api/analyze", json={"content": "Name,Amount\nSales,42"}
        ).json()
        self.assertEqual(bob.get("/api/agents").json(), [])
        self.assertEqual(
            bob.get(f"/api/agents/{agent['id']}/knowledge").status_code, 404
        )
        self.assertEqual(
            bob.post(
                "/api/chat", json={"agent_id": agent["id"], "message": "leave"}
            ).status_code,
            404,
        )
        self.assertEqual(
            bob.get(f"/api/reports/{report['report_id']}/download").status_code, 404
        )
        self.assertEqual(
            bob.put(
                f"/api/agents/{agent['id']}/prompt", json={"instructions": "tamper"}
            ).status_code,
            404,
        )
        doc_id = alice.get(f"/api/agents/{agent['id']}/knowledge").json()["documents"][
            0
        ]["id"]
        self.assertEqual(
            bob.get(f"/api/agents/{agent['id']}/documents/{doc_id}").status_code, 404
        )
        self.assertEqual(
            alice.get(f"/api/agents/{agent['id']}/documents/{doc_id}").content,
            b"Vacation leave is 30 days.",
        )
        exported = load_workbook(io.BytesIO(bob.get("/api/storage/export").content))
        self.assertEqual(exported["agents"].max_row, 1)
        self.assertEqual(exported["users"].max_row, 2)
        self.assertEqual(exported["users"]["B2"].value, "bob@gmail.com")
        exported.close()
        alice.close()
        bob.close()

    def test_google_unverified_identity_is_rejected_and_demo_disabled_when_configured(
        self,
    ):
        google = SimpleNamespace(
            authorize_access_token=AsyncMock(
                return_value={
                    "userinfo": {
                        "sub": "fake",
                        "email": "fake@gmail.com",
                        "email_verified": False,
                    }
                }
            )
        )
        with (
            patch.object(config, "GOOGLE_CONFIGURED", True),
            patch.object(auth, "oauth", SimpleNamespace(google=google)),
        ):
            self.assertEqual(self.client.post("/api/auth/demo").status_code, 403)
            response = self.client.get(
                "/api/auth/google/callback", follow_redirects=False
            )
            self.assertIn("auth_error=google", response.headers["location"])
            self.assertEqual(self.client.get("/api/agents").status_code, 401)

    def test_long_excel_cells_are_lossless_safe_and_rollback_is_atomic(self):
        payload = {"text": "A" * 100_000, "nested": ["unicode café", "=1+1"]}
        with self.store.transaction() as tx:
            tx.insert("reports", {"owner_id": "demo", "body": payload})
            tx.insert("agents", {"owner_id": "demo", "name": '=HYPERLINK("evil")'})
        with ExcelStore(self.store.path).transaction() as tx:
            self.assertEqual(tx.one("reports", owner_id="demo")["body"], payload)
        try:
            with self.store.transaction() as tx:
                tx.insert("runs", {"owner_id": "demo", "agent": "should rollback"})
                raise ValueError("cancel")
        except ValueError:
            pass
        with self.store.transaction() as tx:
            self.assertEqual(tx.all("runs"), [])
        workbook = load_workbook(self.store.path)
        self.assertEqual(workbook["agents"]["C2"].data_type, "s")
        self.assertIn("_long_values", workbook.sheetnames)
        self.assertTrue(set(SCHEMA).issubset(workbook.sheetnames))
        workbook.close()

    def test_agent_files_prompt_edit_and_custom_extension_are_used(self):
        agent = self.client.post(
            "/api/agents",
            json={
                "name": "Payroll Expert",
                "domain": "HR",
                "type": "Chat",
                "instructions": "Original",
            },
        ).json()
        self.assertTrue((self.root / agent["code_path"]).is_file())
        self.assertTrue((self.root / agent["prompt_path"]).is_file())
        self.assertTrue((self.root / agent["knowledge_path"] / "documents").is_dir())
        response = self.client.put(
            f"/api/agents/{agent['id']}/prompt", json={"instructions": "Updated prompt"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            (self.root / agent["prompt_path"]).read_text(), "Updated prompt"
        )
        self.assertEqual(
            self.client.get("/api/agents").json()[0]["instructions"], "Updated prompt"
        )
        code = self.root / agent["code_path"]
        code.write_text(
            'from foundry_agents.chat_agent import ChatAgent\nclass Agent(ChatAgent):\n async def run(self, context):\n  return {"answer": "Custom developer logic", "mode": "Custom", "sources": []}\n'
        )
        response = self.client.post(
            "/api/chat", json={"agent_id": agent["id"], "message": "hi"}
        ).json()
        self.assertEqual(response["answer"], "Custom developer logic")

    def test_sqlite_import_preserves_source_and_is_repeatable(self):
        legacy = self.root / "legacy.db"
        connection = sqlite3.connect(legacy)
        connection.execute(
            "CREATE TABLE agents (id INTEGER PRIMARY KEY,name TEXT,domain TEXT,type TEXT,instructions TEXT)"
        )
        connection.execute(
            "INSERT INTO agents VALUES(7,'Legacy Agent','HR','Chat','Preserved prompt')"
        )
        connection.commit()
        connection.close()
        before = legacy.read_bytes()
        self.assertTrue(migrate_legacy(self.store, legacy))
        self.assertFalse(migrate_legacy(self.store, legacy))
        self.assertEqual(legacy.read_bytes(), before)
        with self.store.transaction() as tx:
            agent = tx.one("agents", id=7, owner_id="demo")
            self.assertEqual(agent["instructions"], "Preserved prompt")
            self.assertTrue((self.root / agent["code_path"]).is_file())

    def test_workbook_save_failure_cleans_new_agent_files_and_restores_prompt(self):
        with patch(
            "platform_core.store.os.replace",
            side_effect=PermissionError("Excel file is open"),
        ):
            response = self.client.post(
                "/api/agents",
                json={"name": "Failed Creation", "domain": "HR", "type": "Chat"},
            )
            self.assertEqual(response.status_code, 503)
        self.assertEqual(self.client.get("/api/agents").json(), [])
        self.assertEqual(list((self.root / "user_agents").rglob("*.py")), [])
        agent = self.client.post(
            "/api/agents",
            json={
                "name": "Existing",
                "domain": "HR",
                "type": "Chat",
                "instructions": "Keep me",
            },
        ).json()
        with patch(
            "platform_core.store.os.replace",
            side_effect=PermissionError("Excel file is open"),
        ):
            self.assertEqual(
                self.client.put(
                    f"/api/agents/{agent['id']}/prompt",
                    json={"instructions": "Should rollback"},
                ).status_code,
                503,
            )
        self.assertEqual((self.root / agent["prompt_path"]).read_text(), "Keep me")

    def test_workflow_dispatches_to_its_capability_module(self):
        response = self.client.post(
            "/api/preview",
            json={"message": "Preview", "template_id": "operations_autopilot"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["simulation"])
        self.assertIn("Human", response.json()["steps"][2])
        self.assertEqual(
            self.client.post(
                "/api/preview",
                json={"message": "Preview", "template_id": "customer_care"},
            ).status_code,
            400,
        )

    def test_concurrent_transactions_keep_unique_ids(self):
        from concurrent.futures import ThreadPoolExecutor

        def write(index):
            with self.store.transaction() as tx:
                return tx.insert("runs", {"owner_id": "demo", "agent": f"run-{index}"})[
                    "id"
                ]

        with ThreadPoolExecutor(max_workers=4) as pool:
            ids = list(pool.map(write, range(12)))
        self.assertEqual(len(set(ids)), 12)
        with self.store.transaction() as tx:
            self.assertEqual(len(tx.all("runs", owner_id="demo")), 12)
