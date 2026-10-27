"""Deterministic CSV/XLSX summaries and safe downloadable reports."""
import csv
import io
import math
import statistics
from datetime import date, datetime
from pathlib import Path

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from knowledge import check_archive

MAX_ROWS = 10000
MAX_COLUMNS = 100
MAX_CELLS = 200000


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    text = '' if value is None else str(value)
    if len(text) > 32767 or ILLEGAL_CHARACTERS_RE.search(text):
        raise ValueError('Cells must have at most 32,767 characters and no unsupported control characters.')
    return text


def parse_table(values):
    try:
        header = next(values)
    except StopIteration:
        raise ValueError('Dataset has no header.')
    header = list(header)
    while header and header[-1] is None:
        header.pop()
    columns = [scalar(value).strip() for value in header]
    if not columns or len(columns) > MAX_COLUMNS or any(not c for c in columns) or len(set(columns)) != len(columns):
        raise ValueError('Use 1–100 nonempty, unique column headers.')
    rows = []
    for values_row in values:
        row = list(values_row)
        if not any(value is not None and str(value).strip() for value in row):
            continue
        if len(row) > len(columns) and any(value is not None and str(value).strip() for value in row[len(columns):]):
            raise ValueError('Rows contain more cells than the header.')
        if len(row) < len(columns):
            raise ValueError('Rows contain fewer cells than the header.')
        rows.append({column: scalar(value).strip() for column, value in zip(columns, row)})
        if len(rows) > MAX_ROWS or len(rows) * len(columns) > MAX_CELLS:
            raise ValueError('Dataset exceeds 10,000 rows or 200,000 cells.')
    if not rows:
        raise ValueError('Dataset contains no data rows.')
    return {'columns': columns, 'data': rows}


def parse_dataset(name, data):
    suffix = Path(name).suffix.lower()
    try:
        if suffix == '.csv':
            text = data.decode('utf-8-sig')
            sample = text[:8192]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            table = parse_table(iter(csv.reader(io.StringIO(text), dialect=dialect, strict=True)))
            return [{'name': 'CSV', **table}]
        if suffix == '.xlsx':
            check_archive(data)
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                if len(workbook.worksheets) > 10:
                    raise ValueError('Workbooks may contain at most 10 worksheets.')
                tables = []
                for sheet in workbook.worksheets:
                    if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                        raise ValueError('Worksheet exceeds 10,000 rows or 100 columns.')
                    values = iter(sheet.iter_rows(values_only=True))
                    # Skip blank worksheets; malformed nonblank sheets still reject the upload.
                    first = next(values, None)
                    if first is None or (not any(v is not None for v in first) and (sheet.max_row or 0) <= 1):
                        continue
                    from itertools import chain
                    tables.append({'name': sheet.title, **parse_table(iter(chain([first], values)))})
                if not tables:
                    raise ValueError('Workbook contains no populated worksheets.')
                if sum(len(t['data']) * len(t['columns']) for t in tables) > MAX_CELLS:
                    raise ValueError('Workbook exceeds 200,000 populated cells.')
                return tables
            finally:
                workbook.close()
        raise ValueError('Upload CSV or Excel XLSX. Convert legacy XLS files to XLSX.')
    except HTTPException:
        raise
    except (ValueError, csv.Error, UnicodeError) as exc:
        raise HTTPException(400, f'Invalid dataset: {exc}') from None
    except Exception:
        raise HTTPException(400, 'Unable to read spreadsheet. Check the file format and integrity.') from None


def summarize(table):
    rows = table['data']
    metrics = []
    quality = []
    for column in table['columns']:
        present = [row[column] for row in rows if row[column] != '']
        numbers = []
        for value in present:
            try:
                number = float(value)
                if not math.isfinite(number):
                    raise ValueError()
                numbers.append(number)
            except ValueError:
                break
        else:
            if numbers:
                try:
                    metrics.append({'column': column, 'count': len(numbers), 'missing': len(rows) - len(present),
                                    'sum': round(math.fsum(numbers), 4), 'average': round(statistics.mean(numbers), 4),
                                    'median': statistics.median(numbers), 'min': min(numbers), 'max': max(numbers)})
                except (OverflowError, ValueError):
                    raise HTTPException(400, f'Numeric values in {column} are too large to summarize.')
        quality.append({'column': column, 'missing': len(rows) - len(present), 'unique': len(set(present))})
    return {'name': table['name'], 'rows': len(rows), 'columns': table['columns'], 'metrics': metrics,
            'quality': quality, 'preview': rows[:8]}


def make_report(filename, tables):
    sheets = [summarize(table) for table in tables]
    return {'filename': filename, 'sheets': sheets,
            'warnings': ['Excel formulas are not recalculated; only cached formula results are read.'] if filename.lower().endswith('.xlsx') else [],
            **{key: value for key, value in sheets[0].items() if key != 'name'}}


def report_workbook(report):
    workbook = Workbook()
    overview = workbook.active
    overview.title = 'Overview'
    overview.append(['Source', report['filename']])
    overview.append(['Generated UTC', report['created']])
    overview.append(['Method', 'Calculated numeric statistics; no generated code execution'])
    for warning in report['warnings']:
        overview.append(['Note', warning])
    metrics = workbook.create_sheet('Statistics')
    metrics.append(['Worksheet', 'Column', 'Count', 'Missing', 'Total', 'Average', 'Median', 'Minimum', 'Maximum'])
    quality = workbook.create_sheet('Data quality')
    quality.append(['Worksheet', 'Column', 'Rows', 'Missing', 'Unique nonempty'])
    for i, sheet in enumerate(report['sheets']):
        overview.append(['Worksheet', sheet['name'], 'Rows', sheet['rows']])
        for m in sheet['metrics']:
            metrics.append([sheet['name'], m['column'], *[m[k] for k in ('count', 'missing', 'sum', 'average', 'median', 'min', 'max')]])
        for q in sheet['quality']:
            quality.append([sheet['name'], q['column'], sheet['rows'], q['missing'], q['unique']])
        preview = workbook.create_sheet(f'Preview {i + 1}')
        preview.append(sheet['columns'])
        for row in sheet['preview']:
            preview.append([row[column] for column in sheet['columns']])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        for row in sheet:
            for cell in row:
                # Uploaded strings must never become executable spreadsheet formulas.
                if cell.data_type == 'f' or isinstance(cell.value, str):
                    cell.data_type = 's'
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
