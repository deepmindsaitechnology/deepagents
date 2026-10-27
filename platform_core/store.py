"""Workbook-backed repository. One lock and atomic replace per transaction; no SQL runtime."""

import io
import json
import os
import tempfile
import uuid
from contextlib import contextmanager
from pathlib import Path

from filelock import FileLock, Timeout
from openpyxl import Workbook, load_workbook

SCHEMA = {
    "users": ["id", "email", "name", "picture", "created"],
    "agents": [
        "id",
        "owner_id",
        "name",
        "domain",
        "type",
        "instructions",
        "source",
        "slug",
        "prompt_path",
        "code_path",
        "knowledge_path",
        "created",
    ],
    "runs": ["id", "owner_id", "agent_id", "agent", "mode", "created"],
    "knowledge_documents": [
        "id",
        "owner_id",
        "agent_id",
        "filename",
        "chunks",
        "file_path",
        "text_path",
        "created",
    ],
    "knowledge_chunks": [
        "id",
        "owner_id",
        "agent_id",
        "source",
        "location",
        "content",
        "embedding",
    ],
    "vector_indexes": ["id", "owner_id", "agent_id", "idf"],
    "datasets": ["id", "owner_id", "agent_id", "filename", "tables", "file_path"],
    "reports": ["id", "owner_id", "agent_id", "body"],
    "auth_events": ["id", "owner_id", "event", "created"],
    "migrations": ["id", "owner_id", "source", "created"],
}


def literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Treat formula-like user input as literal text.


class Transaction:
    def __init__(self, workbook):
        self.workbook = workbook
        self.dirty = False
        self.cache = {}
        for name, columns in SCHEMA.items():
            if name not in workbook.sheetnames:
                sheet = workbook.create_sheet(name)
                sheet.append(columns)
                sheet.freeze_panes = "A2"
                self.dirty = True
        if "_long_values" not in workbook.sheetnames:
            workbook.create_sheet("_long_values").append(["reference", "part", "text"])
            self.dirty = True
        self.long_values = {}
        for ref, part, text in workbook["_long_values"].iter_rows(
            min_row=2, values_only=True
        ):
            self.long_values.setdefault(ref, []).append((part, text or ""))

    def decode(self, value):
        if isinstance(value, str) and value.startswith("@long:"):
            ref = value[6:]
            if ref not in self.long_values:
                raise ValueError("Workbook contains a missing long-value reference.")
            value = "".join(text for _, text in sorted(self.long_values[ref]))
        if isinstance(value, str) and value.startswith("@json:"):
            return json.loads(value[6:])
        if isinstance(value, str) and value.startswith("@str:"):
            return value[5:]
        return value

    def encode(self, value):
        if isinstance(value, (dict, list)):
            value = "@json:" + json.dumps(value, ensure_ascii=True, allow_nan=False)
        elif isinstance(value, str) and (
            value.startswith("@")
            or any(ord(c) < 32 and c not in "\t\n\r" for c in value)
        ):
            value = "@json:" + json.dumps(value, ensure_ascii=True)
        if isinstance(value, str) and len(value) > 28000:
            ref = uuid.uuid4().hex
            pieces = [
                (i // 28000, value[i : i + 28000]) for i in range(0, len(value), 28000)
            ]
            for part, text in pieces:
                sheet = self.workbook["_long_values"]
                sheet.append([ref, part, None])
                literal(sheet.cell(sheet.max_row, 3), text)
            self.long_values[ref] = pieces
            return "@long:" + ref
        return value

    def all(self, table, **where):
        if table not in self.cache:
            columns = SCHEMA[table]
            self.cache[table] = [
                {key: self.decode(value) for key, value in zip(columns, values)}
                for values in self.workbook[table].iter_rows(
                    min_row=2, values_only=True
                )
            ]
        return [
            row
            for row in self.cache[table]
            if all(row.get(key) == value for key, value in where.items())
        ]

    def one(self, table, **where):
        return next(iter(self.all(table, **where)), None)

    def insert(self, table, values):
        values = dict(values)
        if "id" not in values:
            values["id"] = (
                max(
                    (r["id"] for r in self.all(table) if isinstance(r["id"], int)),
                    default=0,
                )
                + 1
            )
        if self.one(table, id=values["id"]):
            raise ValueError(f"Duplicate row ID in {table}.")
        sheet = self.workbook[table]
        sheet.append([None] * len(SCHEMA[table]))
        for column, name in enumerate(SCHEMA[table], 1):
            literal(sheet.cell(sheet.max_row, column), self.encode(values.get(name)))
        self.dirty = True
        result = {name: values.get(name) for name in SCHEMA[table]}
        self.cache[table].append(result)
        return result

    def update(self, table, row_id, values):
        sheet = self.workbook[table]
        for row in sheet.iter_rows(min_row=2):
            if self.decode(row[0].value) == row_id:
                for key, value in values.items():
                    literal(row[SCHEMA[table].index(key)], self.encode(value))
                self.dirty = True
                self.cache.pop(table, None)
                return self.one(table, id=row_id)
        raise KeyError(row_id)


class ExcelStore:
    def __init__(self, path):
        self.path = Path(path)

    @contextmanager
    def transaction(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        try:
            with FileLock(str(self.path) + ".lock", timeout=15):
                if self.path.exists():
                    workbook = load_workbook(self.path)
                else:
                    workbook = Workbook()
                    workbook.remove(workbook.active)
                try:
                    tx = Transaction(workbook)
                    yield tx
                    if tx.dirty:
                        handle, name = tempfile.mkstemp(
                            dir=self.path.parent, suffix=".xlsx"
                        )
                        os.close(handle)
                        try:
                            workbook.save(name)
                            os.replace(name, self.path)
                        finally:
                            Path(name).unlink(missing_ok=True)
                finally:
                    workbook.close()
        except PermissionError as exc:
            raise RuntimeError(
                "Close platform.xlsx in Excel before saving changes, then retry."
            ) from exc
        except Timeout as exc:
            raise RuntimeError(
                "Workbook is busy. Retry in a moment; use one server worker for this prototype."
            ) from exc

    def export(self, owner_id):
        workbook = Workbook()
        workbook.remove(workbook.active)
        target = Transaction(workbook)
        with self.transaction() as source:
            for table in SCHEMA:
                rows = (
                    source.all(table, id=owner_id)
                    if table == "users"
                    else source.all(table, owner_id=owner_id)
                )
                for row in rows:
                    target.insert(table, row)
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()
