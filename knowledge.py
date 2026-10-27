"""Bounded document extraction and local TF-IDF embeddings for a small vector store."""
import io
import json
import math
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT = 300_000
MAX_CHUNKS = 1000
STOP_WORDS = set('a an and are as at be by for from how i in is it of on or our that the their this to was we what when where which who why will with you your'.split())


async def read_upload(file: UploadFile):
    name = (file.filename or 'upload').replace('\\', '/').split('/')[-1][:180]
    data = await file.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(413, f'{name}: maximum file size is 5 MB.')
    if not data:
        raise HTTPException(400, f'{name}: file is empty.')
    return name, data


def check_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 3000 or sum(item.file_size for item in entries) > 30 * 1024 * 1024:
            raise HTTPException(413, 'Office document expands beyond the supported size limit.')


def extract_text(name, data):
    """Return text segments and source locations; files are never executed or written by name."""
    suffix = Path(name).suffix.lower()
    if suffix not in ('.txt', '.md', '.pdf', '.docx'):
        raise HTTPException(400, 'Documents must be UTF-8 TXT/MD, PDF, or Word DOCX. Convert legacy DOC files to DOCX.')
    try:
        if suffix in ('.txt', '.md'):
            segments = [(data.decode('utf-8-sig'), 'text')]
        elif suffix == '.pdf':
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise HTTPException(400, 'Password-protected PDFs are not supported.')
            if len(reader.pages) > 150:
                raise HTTPException(413, 'PDFs may contain at most 150 pages.')
            segments = [(page.extract_text() or '', f'page {i + 1}') for i, page in enumerate(reader.pages)]
        else:
            check_archive(data)
            doc = Document(io.BytesIO(data))
            text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows])
            segments = [(text, 'document')]
        if sum(len(text) for text, _ in segments) > MAX_TEXT:
            raise HTTPException(413, 'Extracted document text exceeds 300,000 characters.')
        segments = [(text.replace('\x00', '').strip(), location) for text, location in segments if text.strip()]
        if not segments:
            raise HTTPException(400, 'No readable text found. Scanned PDFs need OCR before upload.')
        return segments
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, f'{name}: could not read this document. Check encoding or file integrity.') from None


def chunk_document(name, segments):
    chunks = []
    for text, location in segments:
        words = text.split()
        for start in range(0, len(words), 180):
            content = ' '.join(words[start:start + 220])
            if content:
                chunks.append({'source': name, 'location': location, 'text': content})
    return chunks


def tokens(text):
    return [word for word in re.findall(r'\b\w+\b', text.lower(), flags=re.UNICODE) if word not in STOP_WORDS]


def embed(text, idf):
    counts = Counter(tokens(text))
    vector = {word: (1 + math.log(count)) * idf[word] for word, count in counts.items() if word in idf}
    norm = math.sqrt(sum(weight * weight for weight in vector.values()))
    return {word: weight / norm for word, weight in vector.items()} if norm else {}


def index_chunks(db, agent_id, chunks):
    if len(chunks) > MAX_CHUNKS:
        raise HTTPException(413, 'Knowledge exceeds 1,000 text chunks. Use fewer or smaller documents.')
    frequencies = Counter(word for chunk in chunks for word in set(tokens(chunk['text'])))
    idf = {word: math.log((1 + len(chunks)) / (1 + count)) + 1 for word, count in frequencies.items()}
    db.execute('INSERT INTO vector_indexes(agent_id, idf) VALUES(?,?)', (agent_id, json.dumps(idf)))
    for chunk in chunks:
        db.execute('INSERT INTO knowledge_chunks(agent_id, source, location, content, embedding) VALUES(?,?,?,?,?)',
                   (agent_id, chunk['source'], chunk['location'], chunk['text'], json.dumps(embed(chunk['text'], idf))))


def retrieve(db, agent_id, query, limit=4):
    index = db.execute('SELECT idf FROM vector_indexes WHERE agent_id=?', (agent_id,)).fetchone()
    if not index:
        return []
    vector = embed(query, json.loads(index['idf']))
    if not vector:
        return []
    results = []
    for row in db.execute('SELECT id, source, location, content, embedding FROM knowledge_chunks WHERE agent_id=?', (agent_id,)):
        stored = json.loads(row['embedding'])
        score = sum(weight * stored.get(word, 0) for word, weight in vector.items())
        if score > 0.05:
            results.append({'chunk_id': row['id'], 'source': row['source'], 'location': row['location'], 'text': row['content'], 'score': round(score, 4)})
    return sorted(results, key=lambda row: row['score'], reverse=True)[:limit]
